import io
from datetime import date
from xml.sax.saxutils import escape

from fastapi import Response
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4, landscape
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

from app import bitacora

MEDIA = {
    "pdf": "application/pdf",
    "xlsx": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
}

SEXO = {"FEMENINO": "Femenino", "MASCULINO": "Masculino"}

COLUMNAS = [
    ("Fecha", lambda c: c.fechaCirugia.strftime("%d/%m/%Y")),
    ("Hora", lambda c: c.fechaCirugia.strftime("%H:%M")),
    ("Quirófano", lambda c: c.quirofano),
    ("Paciente", lambda c: c.nombrePaciente),
    ("Edad", lambda c: c.edad),
    ("Sexo", lambda c: SEXO.get(c.sexo, c.sexo)),
    ("Procedimiento", lambda c: c.procedimiento),
    ("Diagnóstico", lambda c: c.diagnostico or ""),
    ("Servicios especiales", lambda c: c.serviciosEspeciales or ""),
    ("Cirujano", lambda c: c.cirujano),
    ("Duración (min)", lambda c: c.duracionEstimada),
    ("Estatus", lambda c: c.estatus),
    ("Aseguradora", lambda c: c.aseguradora or "No"),
]
ANCHOS = [48, 30, 54, 76, 26, 48, 78, 86, 82, 70, 42, 64, 60]  # PDF horizontal: suman ≤ 782 pt

COLUMNAS_BITACORA = [
    ("Fecha", lambda b: b.fecha.strftime("%d/%m/%Y")),
    ("Hora", lambda b: b.fecha.strftime("%H:%M:%S")),
    ("Usuario", bitacora.autor),
    ("Acción", lambda b: b.accion),
    ("Entidad", lambda b: b.entidad),
    ("ID", lambda b: b.entidadId or ""),
    ("Detalle", bitacora.resumen),
]
ANCHOS_BITACORA = [52, 46, 140, 85, 55, 30, 370]


def titulo_rango(desde: date, hasta: date) -> str:
    if desde == hasta:
        return f"Cirugías del {desde:%d/%m/%Y}"
    return f"Cirugías del {desde:%d/%m/%Y} al {hasta:%d/%m/%Y}"


def generar_xlsx(filas, titulo: str, columnas=COLUMNAS, hoja: str = "Cirugías") -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = hoja  # Excel no admite / : ? * [ ] en el nombre de la hoja, por eso no se usa el título
    ws.append([n for n, _ in columnas])
    for celda in ws[1]:
        celda.font = Font(bold=True, color="FFFFFF")
        celda.fill = PatternFill("solid", fgColor="1D4ED8")
        celda.alignment = Alignment(vertical="center")
    for fila in filas:
        ws.append([f(fila) for _, f in columnas])
    for renglon in ws.iter_rows(min_row=2):
        for celda in renglon:
            celda.alignment = Alignment(vertical="top", wrap_text=True)  # diagnóstico y servicios pueden ser largos
    for i, col in enumerate(ws.columns, start=1):
        ancho = max(len(str(celda.value or "")) for celda in col)
        ws.column_dimensions[get_column_letter(i)].width = min(ancho + 3, 45)
    ws.freeze_panes = "A2"
    wb.properties.title = titulo
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def generar_pdf(filas, titulo: str, columnas=COLUMNAS, anchos=ANCHOS) -> bytes:
    estilos = getSampleStyleSheet()
    celda = ParagraphStyle("celda", parent=estilos["BodyText"], fontSize=8, leading=10)
    cabecera = ParagraphStyle("cab", parent=celda, textColor=colors.white, fontName="Helvetica-Bold")

    def p(valor, estilo):
        return Paragraph(escape(str(valor)), estilo)  # escape: reportlab interpreta marcado

    datos = [[p(n, cabecera) for n, _ in columnas]]
    datos += [[p(f(fila), celda) for _, f in columnas] for fila in filas]
    tabla = Table(datos, repeatRows=1, colWidths=anchos)
    tabla.setStyle(TableStyle([
        ("LEFTPADDING", (0, 0), (-1, -1), 3),  # 13 columnas: el padding por defecto (6) parte las palabras
        ("RIGHTPADDING", (0, 0), (-1, -1), 3),
        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#1D4ED8")),
        ("GRID", (0, 0), (-1, -1), 0.4, colors.HexColor("#CBD5E1")),
        ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
        ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#F4F6FA")]),
    ]))
    buf = io.BytesIO()
    doc = SimpleDocTemplate(buf, pagesize=landscape(A4), title=titulo,
                            leftMargin=30, rightMargin=30, topMargin=30, bottomMargin=30)
    doc.build([Paragraph(escape(titulo), estilos["Title"]), Spacer(1, 10), tabla])
    return buf.getvalue()


def respuesta_archivo(filas, formato: str, titulo: str, nombre: str,
                      columnas=COLUMNAS, anchos=ANCHOS, hoja: str = "Cirugías") -> Response:
    contenido = (generar_pdf(filas, titulo, columnas, anchos) if formato == "pdf"
                 else generar_xlsx(filas, titulo, columnas, hoja))
    return Response(
        content=contenido,
        media_type=MEDIA[formato],
        headers={"Content-Disposition": f'attachment; filename="{nombre}.{formato}"'},
    )
